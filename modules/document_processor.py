"""
Document Processor Module
Handles document upload validation and text extraction from PDF, DOCX, and TXT files.
Supports multi-page extraction, table extraction for Word documents, encoding detection,
and fallback strategies for corrupted or encrypted files.
"""

import os
from pathlib import Path
from typing import Dict, Any, Tuple


class DocumentProcessingError(Exception):
    """Custom exception raised when document extraction or validation fails."""
    pass


class DocumentProcessor:
    """Extracts clean text and metadata from supported legal document formats."""

    ALLOWED_EXTENSIONS = {'pdf', 'docx', 'txt'}

    def __init__(self, upload_folder: str = None):
        self.upload_folder = upload_folder

    @staticmethod
    def allowed_file(filename: str) -> bool:
        """Check if the filename has an allowed extension."""
        if not filename or '.' not in filename:
            return False
        ext = filename.rsplit('.', 1)[1].lower()
        return ext in DocumentProcessor.ALLOWED_EXTENSIONS

    @classmethod
    def get_extension(cls, filename: str) -> str:
        """Return lowercase extension without dot."""
        if '.' in filename:
            return filename.rsplit('.', 1)[1].lower()
        return ''

    def extract_text(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """
        Extract text and basic metadata from the given file path.
        
        Args:
            file_path: Absolute or relative path to the file.
            
        Returns:
            Tuple of (extracted_text, metadata_dict)
            
        Raises:
            DocumentProcessingError if extraction fails or document is unreadable/empty.
        """
        if not os.path.exists(file_path):
            raise DocumentProcessingError(f"File not found on server: {file_path}")

        file_size = os.path.getsize(file_path)
        if file_size == 0:
            raise DocumentProcessingError("The uploaded file is empty (0 bytes). Please upload a valid document.")

        filename = os.path.basename(file_path)
        extension = self.get_extension(filename)

        if extension not in self.ALLOWED_EXTENSIONS:
            raise DocumentProcessingError(
                f"Unsupported file format '.{extension}'. Allowed formats are: {', '.join(self.ALLOWED_EXTENSIONS)}"
            )

        try:
            if extension == 'pdf':
                text, extra_meta = self._extract_from_pdf(file_path)
            elif extension == 'docx':
                text, extra_meta = self._extract_from_docx(file_path)
            elif extension == 'txt':
                text, extra_meta = self._extract_from_txt(file_path)
            else:
                raise DocumentProcessingError(f"Handler for extension '.{extension}' not implemented.")

        except DocumentProcessingError:
            raise
        except Exception as e:
            raise DocumentProcessingError(f"Failed to process '{filename}': {str(e)}") from e

        cleaned_text = text.strip()
        if not cleaned_text:
            raise DocumentProcessingError(
                "No readable text could be extracted from this document. "
                "If this is a scanned PDF image or an empty document, please upload a text-searchable document."
            )

        metadata = {
            'filename': filename,
            'file_extension': extension,
            'file_size_bytes': file_size,
            'file_size_formatted': self._format_file_size(file_size),
            'char_count': len(cleaned_text),
            'word_count': len(cleaned_text.split()),
            **extra_meta
        }

        return cleaned_text, metadata

    def _extract_from_pdf(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract text from PDF using PyMuPDF (fitz) with fallback to pypdf."""
        extracted_pages = []
        page_count = 0

        # Primary extraction using PyMuPDF
        pymupdf_success = False
        try:
            import pymupdf  # PyMuPDF
            doc = pymupdf.open(file_path)
            if doc.is_encrypted:
                raise DocumentProcessingError("This PDF is password-protected or encrypted. Please decrypt it first.")

            page_count = len(doc)
            for page_index in range(page_count):
                page = doc[page_index]
                page_text = page.get_text("text")
                if page_text and page_text.strip():
                    extracted_pages.append(page_text.strip())
            doc.close()
            pymupdf_success = True
        except ImportError:
            pymupdf_success = False
        except DocumentProcessingError:
            raise
        except Exception:
            pymupdf_success = False

        # Fallback extraction using pypdf if PyMuPDF wasn't used or failed
        if not pymupdf_success or not extracted_pages:
            try:
                import pypdf
                reader = pypdf.PdfReader(file_path)
                if reader.is_encrypted:
                    raise DocumentProcessingError("This PDF is password-protected or encrypted. Please decrypt it first.")
                
                page_count = len(reader.pages)
                extracted_pages = []
                for page in reader.pages:
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        extracted_pages.append(page_text.strip())
            except DocumentProcessingError:
                raise
            except Exception as e:
                if not extracted_pages:
                    raise DocumentProcessingError(f"Could not parse PDF content: {str(e)}")

        full_text = "\n\n".join(extracted_pages)
        return full_text, {
            'page_count': page_count,
            'readable_pages': len(extracted_pages)
        }

    def _extract_from_docx(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract text from Word .docx file including paragraphs and tables."""
        try:
            import docx
            doc = docx.Document(file_path)
        except Exception as e:
            raise DocumentProcessingError(f"Invalid or corrupted Word (.docx) document: {str(e)}")

        paragraphs = []
        for p in doc.paragraphs:
            p_text = p.text.strip()
            if p_text:
                paragraphs.append(p_text)

        # Also extract table text (common for legal agreement terms/schedules)
        table_cells = []
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    # Remove adjacent duplicates from merged cells
                    deduped = []
                    for item in row_text:
                        if not deduped or item != deduped[-1]:
                            deduped.append(item)
                    table_cells.append(" | ".join(deduped))

        combined = paragraphs + table_cells
        full_text = "\n\n".join(combined)

        return full_text, {
            'paragraph_count': len(paragraphs),
            'table_count': len(doc.tables)
        }

    def _extract_from_txt(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract text from plain text file trying common text encodings."""
        encodings_to_try = ['utf-8', 'utf-8-sig', 'windows-1252', 'latin-1', 'cp1252']
        content = None
        used_encoding = None

        with open(file_path, 'rb') as f:
            raw_bytes = f.read()

        for enc in encodings_to_try:
            try:
                content = raw_bytes.decode(enc)
                used_encoding = enc
                break
            except (UnicodeDecodeError, LookupError):
                continue

        if content is None:
            # Final fallback with replacement
            content = raw_bytes.decode('utf-8', errors='replace')
            used_encoding = 'utf-8 (lossy)'

        # Normalize line endings
        content = content.replace('\r\n', '\n').replace('\r', '\n')
        line_count = len([line for line in content.split('\n') if line.strip()])

        return content, {
            'encoding': used_encoding,
            'line_count': line_count
        }

    @staticmethod
    def _format_file_size(size_bytes: int) -> str:
        """Format byte size into human readable string (KB, MB)."""
        if size_bytes < 1024:
            return f"{size_bytes} B"
        elif size_bytes < 1024 * 1024:
            return f"{size_bytes / 1024:.1f} KB"
        else:
            return f"{size_bytes / (1024 * 1024):.2f} MB"
